import io

from docx import Document

from app.document_parser import extract_paragraphs_docx
from app.docx_builder import build_translated_docx, translations_by_paragraph


def _build_sample_docx() -> bytes:
    document = Document()
    document.add_heading("Original Title", level=1)
    document.add_paragraph("Plain paragraph text.")
    bold_paragraph = document.add_paragraph()
    bold_paragraph.add_run("Bold sentence.").bold = True
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Cell one"
    table.cell(0, 1).text = "Cell two"

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_build_translated_docx_replaces_text_and_keeps_paragraph_style():
    original = _build_sample_docx()
    paragraphs = extract_paragraphs_docx(original)
    translations = {p: f"[TR] {p}" for p in paragraphs}

    rebuilt = Document(io.BytesIO(build_translated_docx(original, translations)))
    texts = [p.text for p in rebuilt.paragraphs if p.text.strip()]

    assert "[TR] Original Title" in texts
    assert "[TR] Plain paragraph text." in texts
    title_paragraph = next(p for p in rebuilt.paragraphs if p.text == "[TR] Original Title")
    assert title_paragraph.style.name.startswith("Heading")


def test_build_translated_docx_keeps_bold_formatting_of_first_run():
    original = _build_sample_docx()
    paragraphs = extract_paragraphs_docx(original)
    translations = {p: f"[TR] {p}" for p in paragraphs}

    rebuilt = Document(io.BytesIO(build_translated_docx(original, translations)))
    bold_paragraph = next(p for p in rebuilt.paragraphs if p.text == "[TR] Bold sentence.")

    assert bold_paragraph.runs[0].bold is True


def test_build_translated_docx_translates_table_cells():
    original = _build_sample_docx()
    paragraphs = extract_paragraphs_docx(original)
    assert "Cell one" in paragraphs
    assert "Cell two" in paragraphs
    translations = {p: f"[TR] {p}" for p in paragraphs}

    rebuilt = Document(io.BytesIO(build_translated_docx(original, translations)))
    cell_texts = [cell.text for table in rebuilt.tables for row in table.rows for cell in row.cells]

    assert "[TR] Cell one" in cell_texts
    assert "[TR] Cell two" in cell_texts


def test_build_translated_docx_leaves_unmapped_paragraphs_untouched():
    original = _build_sample_docx()

    rebuilt = Document(io.BytesIO(build_translated_docx(original, {})))
    texts = [p.text for p in rebuilt.paragraphs if p.text.strip()]

    assert "Original Title" in texts
    assert "Plain paragraph text." in texts


def test_translations_by_paragraph_joins_sentences_per_paragraph():
    paragraphs = ["First para.", "Second para has two sentences."]
    segments = [
        {"paragraph_index": 0, "translation": "First translated."},
        {"paragraph_index": 1, "translation": "Second translated part one."},
        {"paragraph_index": 1, "translation": "Part two."},
    ]

    result = translations_by_paragraph(paragraphs, segments)

    assert result == {
        "First para.": "First translated.",
        "Second para has two sentences.": "Second translated part one. Part two.",
    }


def test_translations_by_paragraph_ignores_out_of_range_index():
    result = translations_by_paragraph(["Only paragraph."], [{"paragraph_index": 5, "translation": "orphan"}])
    assert result == {}
