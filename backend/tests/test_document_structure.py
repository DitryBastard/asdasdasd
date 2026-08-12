import io

import pytest
from docx import Document
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.document_parser import extract_structure, extract_structure_docx, extract_structure_pdf


def _build_sample_docx() -> bytes:
    document = Document()
    document.add_paragraph("Intro paragraph about the material.")
    table = document.add_table(rows=2, cols=3)
    values = [["Fe", "Ti", "Mo"], ["178.29", "0.5", "I"]]
    for row_idx, row_values in enumerate(values):
        for col_idx, value in enumerate(row_values):
            table.cell(row_idx, col_idx).text = value
    document.add_paragraph("Closing paragraph after the table.")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _build_sample_pdf() -> bytes:
    # Intro and closing text live on separate pages (rather than the same
    # page with just a Spacer between them) because pdfplumber's basic text
    # extraction doesn't reliably turn a visual gap into a blank line in the
    # extracted text - relying on that would make this test flaky. Page
    # boundaries, by contrast, always keep text apart: each page is
    # extracted and paragraph-split independently.
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer)
    styles = getSampleStyleSheet()
    table = Table([["Fe", "Ti", "Mo"], ["178.29", "0.5", "I"]])
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
    elements = [
        Paragraph("Intro paragraph about the material.", styles["Normal"]),
        PageBreak(),
        table,
        Spacer(1, 12),
        Paragraph("Closing paragraph after the table.", styles["Normal"]),
    ]
    doc.build(elements)
    return buffer.getvalue()


def test_extract_structure_docx_keeps_table_separate_from_paragraphs():
    paragraphs, tables = extract_structure_docx(_build_sample_docx())

    assert paragraphs == ["Intro paragraph about the material.", "Closing paragraph after the table."]
    assert tables == [[["Fe", "Ti", "Mo"], ["178.29", "0.5", "I"]]]


def test_extract_structure_docx_rejects_malformed_file():
    with pytest.raises(ValueError, match="docx"):
        extract_structure_docx(b"not a docx file")


def test_extract_structure_pdf_detects_table_and_prose_separately():
    paragraphs, tables = extract_structure_pdf(_build_sample_pdf())

    assert any("Intro paragraph" in p for p in paragraphs)
    assert any("Closing paragraph" in p for p in paragraphs)
    # The table's own cell values must not also leak into the prose stream -
    # otherwise cells get translated/aligned twice, once as loose fragments
    # and once as structured table data.
    assert not any("Fe" in p and "Ti" in p for p in paragraphs)
    assert len(tables) == 1
    assert tables[0][0] == ["Fe", "Ti", "Mo"]
    assert tables[0][1] == ["178.29", "0.5", "I"]


def test_extract_structure_dispatches_by_extension():
    docx_paragraphs, docx_tables = extract_structure("original.docx", _build_sample_docx())
    pdf_paragraphs, pdf_tables = extract_structure("original.pdf", _build_sample_pdf())

    assert docx_tables and pdf_tables
    assert docx_paragraphs[0] == pdf_paragraphs[0] == "Intro paragraph about the material."


def test_extract_structure_txt_has_no_tables():
    paragraphs, tables = extract_structure("notes.txt", "First.\n\nSecond.".encode("utf-8"))
    assert paragraphs == ["First.", "Second."]
    assert tables == []


def test_extract_structure_rejects_unsupported_extension():
    with pytest.raises(ValueError, match="Неподдерживаемый"):
        extract_structure("notes.rtf", b"whatever")
